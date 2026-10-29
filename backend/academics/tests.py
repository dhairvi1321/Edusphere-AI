from datetime import date, timedelta
import json
import os
import tempfile
from pathlib import Path
from io import BytesIO, StringIO
from uuid import uuid4
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase
from django.test.utils import override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from rest_framework.test import APIClient
from accounts.models import User
from accounts.views import make_tokens
from mongoengine.connection import get_db
from mongoengine.errors import DoesNotExist
from pymongo.errors import ServerSelectionTimeoutError

from .models import Assignment, QuizAttempt, QuizQuestion, StudyPlan, Subject, SyllabusTopic, TopicQuiz
from .views import (
    AcademicAssistantAsk,
    GenerateStudyPlan,
    StudyInsights,
    StudyPlanListCreate,
    StudyPlanDetail,
    TopicDetail,
    TopicQuizList,
    TopicQuizDetail,
    TopicQuizAttempts,
    _build_session_pool,
    _priority_score,
    _schedule_sessions,
    _topic_duration,
    _topic_priority_key,
)


class StudyPlannerSchedulingTests(SimpleTestCase):
    def setUp(self):
        self.start = date(2026, 10, 29)
        self.end = date(2026, 10, 31)

    @staticmethod
    def topic(topic_id, difficulty=1):
        return {
            'subject': 'Math',
            'subject_id': 'subject-1',
            'topic': f'Topic {topic_id}',
            'topic_id': topic_id,
            '_difficulty_key': difficulty,
        }

    def test_sessions_obey_daily_budget_and_dates(self):
        pool = [self.topic('hard', 0), self.topic('medium', 1), self.topic('easy', 2)]
        sessions = _schedule_sessions(pool, self.start, self.end, 1)

        self.assertEqual(len({s['topic_id'] for s in sessions}), len(sessions))
        self.assertTrue(all(self.start.isoformat() <= s['date'] <= self.end.isoformat() for s in sessions))
        self.assertTrue(all(s['duration_minutes'] <= 60 for s in sessions))
        for session_date in {s['date'] for s in sessions}:
            day_total = sum(s['duration_minutes'] for s in sessions if s['date'] == session_date)
            self.assertLessEqual(day_total, 60)

    def test_duplicate_topic_ids_are_only_scheduled_once(self):
        topic = self.topic('same')
        sessions = _schedule_sessions([topic, dict(topic)], self.start, self.end, 2)
        self.assertEqual([s['topic_id'] for s in sessions], ['same'])

    def test_insufficient_time_leaves_topics_unscheduled_without_overbooking(self):
        pool = [self.topic(str(index), 2) for index in range(8)]
        sessions = _schedule_sessions(pool, self.start, self.start, 1)

        self.assertEqual(len(sessions), 3)
        self.assertLess(len(sessions), len(pool))
        self.assertLessEqual(sum(s['duration_minutes'] for s in sessions), 60)

    def test_invalid_capacity_and_date_range_return_empty(self):
        pool = [self.topic('one')]
        self.assertEqual(_schedule_sessions([], self.start, self.end, 1), [])
        self.assertEqual(_schedule_sessions(pool, self.start, self.end, 0), [])
        self.assertEqual(_schedule_sessions(pool, self.end, self.start, 1), [])
        self.assertEqual(_schedule_sessions(pool, self.start, self.end, True), [])

    def test_unknown_difficulty_uses_safe_medium_duration(self):
        self.assertEqual(_topic_duration({'_difficulty_key': 99}), 30)

    @patch('academics.views.Assignment.objects')
    @patch('academics.views.Exam.objects')
    @patch('academics.views.Subject.objects')
    def test_missing_or_invalid_deadlines_and_difficulty_use_safe_defaults(
        self, subjects, exams, assignments
    ):
        completed = SimpleNamespace(
            id='completed', title='Done', is_completed=True, difficulty='hard'
        )
        pending = SimpleNamespace(
            id='pending', title='Pending', is_completed=False, difficulty='unexpected'
        )
        subject = SimpleNamespace(id='subject-1', name='Math', topics=[completed, pending])
        subjects.return_value = [subject]
        exams.return_value = [SimpleNamespace(subject_id='subject-1', exam_date='invalid-date')]
        assignments.return_value = [SimpleNamespace(
            subject_id='subject-1', is_completed=False, due_date='invalid-date', priority='unexpected'
        )]

        pool = _build_session_pool('student-1')

        self.assertEqual([item['topic_id'] for item in pool], ['pending'])
        self.assertEqual(pool[0]['_difficulty_key'], 1)
        self.assertIsNone(pool[0]['deadline_days'])
        self.assertAlmostEqual(pool[0]['_priority_score'], 0.4166667)

    @patch('academics.views.Assignment.objects')
    @patch('academics.views.Exam.objects')
    @patch('academics.views.Subject.objects')
    def test_nearest_relevant_deadline_drives_urgency(self, subjects, exams, assignments):
        topic = SimpleNamespace(
            id='topic-1', title='Soon', is_completed=False, difficulty='medium',
            status='not_started', revision_required=False,
        )
        subject = SimpleNamespace(id='subject-1', name='Math', topics=[topic])
        subjects.return_value = [subject]
        exams.return_value = [SimpleNamespace(
            subject_id='subject-1', exam_date=(date.today() + timedelta(days=1)).isoformat()
        )]
        assignments.return_value = [SimpleNamespace(
            subject_id='subject-1', is_completed=False,
            due_date=(date.today() + timedelta(days=5)).isoformat(),
        )]

        pool = _build_session_pool('student-1')

        self.assertEqual(pool[0]['deadline_days'], 1)
        self.assertAlmostEqual(pool[0]['_priority_score'], 0.6166667)

    def test_documented_worked_example_scores_and_order(self):
        topic_a = _priority_score(5, 1.0, 'not_started', False)
        topic_b = _priority_score(2, 2 / 3, 'in_progress', True)
        topic_c = _priority_score(10, 1 / 3, 'not_started', False)

        self.assertAlmostEqual(topic_a, 0.567, places=3)
        self.assertAlmostEqual(topic_b, 0.525, places=3)
        self.assertAlmostEqual(topic_c, 0.370, places=3)
        items = [
            {'topic': 'Topic C', 'subject': 'Math', 'topic_id': 'c', '_priority_score': topic_c},
            {'topic': 'Topic B', 'subject': 'Math', 'topic_id': 'b', '_priority_score': topic_b},
            {'topic': 'Topic A', 'subject': 'Math', 'topic_id': 'a', '_priority_score': topic_a},
        ]
        self.assertEqual([item['topic'] for item in sorted(items, key=_topic_priority_key)], [
            'Topic A', 'Topic B', 'Topic C'
        ])

    def test_revision_flag_increases_score_and_equal_scores_break_by_name(self):
        base_score = _priority_score(4, 2 / 3, 'in_progress', False)
        revision_score = _priority_score(4, 2 / 3, 'in_progress', True)
        self.assertAlmostEqual(revision_score - base_score, 0.10)

        common = {'_priority_score': 0.5, 'subject': 'Math', 'topic_id': 'same'}
        items = [
            {**common, 'topic': 'Zeta'},
            {**common, 'topic': 'Alpha'},
        ]
        self.assertEqual([item['topic'] for item in sorted(items, key=_topic_priority_key)], ['Alpha', 'Zeta'])

    def test_missing_deadline_has_zero_urgency_and_overdue_deadline_is_clamped(self):
        no_deadline = _priority_score(None, 2 / 3, 'not_started', False)
        invalid_deadline = _priority_score('invalid', 2 / 3, 'not_started', False)
        overdue = _priority_score(-10, 2 / 3, 'not_started', False)

        self.assertAlmostEqual(no_deadline, invalid_deadline)
        self.assertGreater(overdue, no_deadline)


class StudySessionCompletionTests(SimpleTestCase):
    def setUp(self):
        self.topic = SimpleNamespace(id='topic-1', is_completed=False)
        self.other_topic = SimpleNamespace(id='topic-2', is_completed=False)
        self.subject = SimpleNamespace(topics=[self.topic, self.other_topic], save=MagicMock())
        self.plan = SimpleNamespace(
            id='plan-1',
            title='Study',
            start_date='2026-10-29',
            end_date='2026-10-29',
            daily_hours='2',
            created_at=date(2026, 10, 29),
            sessions=[{
                'session_id': 'session-1',
                'subject_id': 'subject-1',
                'topic_id': 'topic-1',
                'is_completed': False,
            }],
            save=MagicMock(),
        )
        self.request = SimpleNamespace(
            user=SimpleNamespace(id='student-1'),
            data={'session_id': 'session-1', 'is_completed': True},
        )

    @patch('academics.views.Subject.objects')
    def test_completing_session_persists_topic_and_session_for_owner(self, subjects):
        subjects.return_value.first.return_value = self.subject
        with patch.object(StudyPlanDetail, '_get', return_value=self.plan):
            response = StudyPlanDetail().patch(self.request, 'plan-1')

        self.assertEqual(response.status_code, 200)
        self.assertTrue(self.topic.is_completed)
        self.assertFalse(self.other_topic.is_completed)
        self.assertTrue(self.plan.sessions[0]['is_completed'])
        self.subject.save.assert_called_once_with()
        self.plan.save.assert_called_once_with()
        subjects.assert_called_once_with(user_id='student-1', id='subject-1')

    @patch('academics.views.Subject.objects')
    def test_repeating_completion_is_idempotent(self, subjects):
        subjects.return_value.first.return_value = self.subject
        with patch.object(StudyPlanDetail, '_get', return_value=self.plan):
            first = StudyPlanDetail().patch(self.request, 'plan-1')
            second = StudyPlanDetail().patch(self.request, 'plan-1')

        self.assertEqual((first.status_code, second.status_code), (200, 200))
        self.assertTrue(self.topic.is_completed)
        self.assertTrue(self.plan.sessions[0]['is_completed'])

    def test_non_boolean_completion_is_rejected(self):
        self.request.data['is_completed'] = 'false'
        with patch.object(StudyPlanDetail, '_get', return_value=self.plan):
            response = StudyPlanDetail().patch(self.request, 'plan-1')

        self.assertEqual(response.status_code, 400)
        self.plan.save.assert_not_called()


class TopicRevisionStatusTests(SimpleTestCase):
    def test_topic_edit_persists_status_and_revision_flag(self):
        topic = SimpleNamespace(
            id='topic-1', title='Topic', difficulty='medium', is_completed=False,
            status='not_started', revision_required=False,
        )
        subject = SimpleNamespace(id='subject-1', save=MagicMock())
        request = SimpleNamespace(
            user=SimpleNamespace(id='student-1'),
            data={'status': 'in_progress', 'revision_required': True},
        )
        with patch.object(TopicDetail, '_find', return_value=(subject, topic)):
            response = TopicDetail().patch(request, 'topic-1')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(topic.status, 'in_progress')
        self.assertFalse(topic.is_completed)
        self.assertTrue(topic.revision_required)
        subject.save.assert_called_once_with()


class AcademicAssistantTests(SimpleTestCase):
    def setUp(self):
        self.request = SimpleNamespace(
            user=SimpleNamespace(id='student-1'),
            data={'question': 'What assignments are pending?'},
        )

    @patch('academics.views.StudyPlan.objects')
    @patch('academics.views.Exam.objects')
    @patch('academics.views.Assignment.objects')
    @patch('academics.views.Subject.objects')
    def test_pending_tasks_uses_only_authenticated_student_records(self, subjects, assignments, exams, plans):
        subjects.return_value = [SimpleNamespace(id='subject-1', name='Math', topics=[])]
        assignments.return_value = [
            SimpleNamespace(
                title='Algebra worksheet', subject_id='subject-1', due_date='2026-11-02',
                priority='high', is_completed=False,
            ),
            SimpleNamespace(
                title='Already submitted', subject_id='subject-1', due_date='2026-10-20',
                priority='medium', is_completed=True,
            ),
        ]
        exams.return_value = []
        plans.return_value = []

        response = AcademicAssistantAsk().post(self.request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['intent'], 'assignments')
        self.assertIn('Algebra worksheet', response.data['answer'])
        self.assertIn('Math', response.data['answer'])
        self.assertNotIn('Already submitted', response.data['answer'])
        subjects.assert_called_once_with(user_id='student-1')
        assignments.assert_called_once_with(user_id='student-1')
        exams.assert_called_once_with(user_id='student-1')
        plans.assert_called_once_with(user_id='student-1')

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[])
    def test_empty_student_data_returns_clear_answer(self, subjects, assignments, exams, plans):
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.status_code, 200)
        self.assertIn('no pending assignments', response.data['answer'])

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[SimpleNamespace(
        id='subject-1', name='Physics', topics=[
            SimpleNamespace(id='done', title='Done', difficulty='easy', is_completed=True, status='completed', revision_required=False),
            SimpleNamespace(id='pending', title='Pending topic', difficulty='hard', is_completed=False, status='in_progress', revision_required=False),
        ],
    )])
    def test_progress_answer_uses_saved_topic_completion(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'How much syllabus have I completed?'
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'progress')
        self.assertIn('1 of 2 topics (50%)', response.data['answer'])
        self.assertIn('Physics: 1/2', response.data['answer'])

    @patch('academics.views.StudyPlan.objects', return_value=[SimpleNamespace(
        title='Current plan', start_date=date.today().isoformat(), end_date=date.today().isoformat(),
        daily_hours='2', created_at=date.today(), sessions=[
            {'date': date.today().isoformat(), 'topic': 'Arrays', 'subject': 'CS', 'topic_id': 'arrays', 'is_completed': False},
            {'date': date.today().isoformat(), 'topic': 'Old completed session', 'is_completed': True},
        ],
    )])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[])
    def test_schedule_answer_skips_completed_sessions(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'What is in my study schedule?'
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'schedule')
        self.assertIn('Arrays', response.data['answer'])
        self.assertNotIn('Old completed session', response.data['answer'])

    @patch('academics.views.StudyPlan.objects', return_value=[SimpleNamespace(
        title='Midterm plan', start_date='2026-11-01', end_date='2026-11-05', daily_hours='2',
        created_at=date.today(), sessions=[],
    )])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[])
    def test_schedule_change_question_explains_supported_edits_without_mutating(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'what can i change in my schedule'
        saved_plan = plans.return_value[0]
        saved_plan.save = MagicMock()
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'schedule_change')
        self.assertIn('Reschedule option', response.data['answer'])
        self.assertIn('start date, end date, or daily study hours', response.data['answer'])
        self.assertIn('I have not changed your plan', response.data['answer'])
        saved_plan.save.assert_not_called()

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[SimpleNamespace(
        title='Final exam', subject_id='subject-1', exam_date=(date.today() + timedelta(days=3)).isoformat()
    )])
    @patch('academics.views.Assignment.objects', return_value=[SimpleNamespace(
        title='Lab report', subject_id='subject-1', due_date=(date.today() + timedelta(days=1)).isoformat(),
        priority='high', is_completed=False,
    )])
    @patch('academics.views.Subject.objects', return_value=[SimpleNamespace(id='subject-1', name='Software', topics=[
        SimpleNamespace(id='topic-1', title='Testing', difficulty='hard', is_completed=False,
                        status='not_started', revision_required=True),
    ])])
    def test_priority_wording_compares_real_deadlines_and_topic_score(self, subjects, assignments, exams, plans):
        for question in ('What should I prefer first?', 'what should i study first', 'Which assignment is most urgent?'):
            self.request.data['question'] = question
            response = AcademicAssistantAsk().post(self.request)
            self.assertEqual(response.data['intent'], 'priority')
            self.assertIn('Start with assignment “Lab report”', response.data['answer'])
            self.assertIn('Final exam', response.data['answer'])
            self.assertIn('Testing', response.data['answer'])
            self.assertIn('documented Study Planner score', response.data['answer'])
        self.request.data['question'] = 'what can i change in my schedule'
        schedule_response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(schedule_response.data['intent'], 'schedule_change')
        self.assertNotEqual(schedule_response.data['answer'], response.data['answer'])

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[])
    def test_spelling_and_punctuation_variations_match_schedule_intent(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'WhAt can I change in my schedual?!'
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'schedule_change')
        self.assertIn('no saved study plan', response.data['answer'])

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[])
    def test_unsupported_question_gets_clarification_instead_of_overview(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'Tell me a joke.'
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'unsupported')
        self.assertIn('could not match', response.data['answer'])
        self.assertNotIn('Your overview:', response.data['answer'])

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[SimpleNamespace(
        title='Untimed assignment', subject_id=None, due_date=None, priority='medium', is_completed=False
    )])
    @patch('academics.views.Subject.objects', return_value=[])
    def test_missing_assignment_deadline_is_not_invented(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'What is my next assignment deadline?'
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'assignments')
        self.assertIn('due date unavailable', response.data['answer'])
        self.assertIn('cannot rank them by deadline', response.data['answer'])

    def test_empty_or_oversized_questions_are_rejected(self):
        self.request.data['question'] = '  '
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.status_code, 400)
        self.request.data['question'] = 'x' * 501
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.status_code, 400)

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[
        SimpleNamespace(title='Final', subject_id=None, exam_date=(date.today() + timedelta(days=3)).isoformat()),
        SimpleNamespace(title='Past exam', subject_id=None, exam_date=(date.today() - timedelta(days=1)).isoformat()),
        SimpleNamespace(title='Invalid date', subject_id=None, exam_date='not-a-date'),
    ])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[])
    def test_exam_answer_lists_only_upcoming_valid_dates(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'Which exam is coming up next?'
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'exams')
        self.assertIn('Final', response.data['answer'])
        self.assertNotIn('Past exam', response.data['answer'])
        self.assertNotIn('Invalid date', response.data['answer'])

    @patch('academics.views.StudyPlan.objects', return_value=[])
    @patch('academics.views.Exam.objects', return_value=[])
    @patch('academics.views.Assignment.objects', return_value=[])
    @patch('academics.views.Subject.objects', return_value=[SimpleNamespace(
        id='subject-1', name='History', topics=[SimpleNamespace(is_completed=True, status='completed')]
    )])
    def test_general_question_returns_data_backed_overview(self, subjects, assignments, exams, plans):
        self.request.data['question'] = 'Can I see my overview?'
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.data['intent'], 'overview')
        self.assertIn('1 subject(s)', response.data['answer'])
        self.assertIn('1 of 1 topics complete', response.data['answer'])

    @patch('academics.views.Subject.objects', side_effect=ServerSelectionTimeoutError('database unavailable'))
    def test_database_unavailable_returns_service_unavailable(self, subjects):
        response = AcademicAssistantAsk().post(self.request)
        self.assertEqual(response.status_code, 503)


@override_settings(QUIZ_AI_GENERATION_ENABLED=True)
class TopicQuizAPITests(SimpleTestCase):
    def setUp(self):
        self.user_id = 'student-quiz-test'
        self.subject = SimpleNamespace(id='subject-1', name='Software Engineering', topics=[], syllabus_files=[{'filename': 'source.txt', 'name': 'source.txt'}])
        self.topic = SimpleNamespace(id='topic-1', title='Software Testing')
        self.subject.topics = [self.topic]
        self.quiz = SimpleNamespace(
            id='quiz-1', title='Testing Basics', source='author_created', subject_id='subject-1', topic_id='topic-1',
            questions=[
                QuizQuestion(question_id='q1', text='What is a test case?', options=['A check', 'A program'], correct_index=0, explanation='A test case defines input and expected output.'),
                QuizQuestion(question_id='q2', text='Which is a test level?', options=['Unit', 'Compile'], correct_index=0, explanation='Unit testing is a test level.'),
            ],
        )
        self.request = SimpleNamespace(
            user=SimpleNamespace(id=self.user_id),
            data={},
            query_params={'subject_id': 'subject-1', 'topic_id': 'topic-1'},
        )

    def test_quiz_list_is_scoped_to_owned_subject_and_topic(self):
        with patch('academics.views.Subject.objects', return_value=[self.subject]), \
                patch('academics.views.TopicQuiz.objects') as quizzes:
            quizzes.return_value.order_by.return_value = [self.quiz]
            response = TopicQuizList().get(self.request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]['topic_id'], 'topic-1')
        self.assertEqual(response.data[0]['subject_name'], 'Software Engineering')
        quizzes.assert_called_once_with(user_id=self.user_id, subject_id='subject-1')

    def test_topic_list_does_not_return_question_keys(self):
        with patch('academics.views._owned_topic_quiz', return_value=(self.quiz, self.subject, self.topic)):
            response = TopicQuizDetail().get(self.request, 'quiz-1')

        self.assertEqual(response.status_code, 200)
        question = response.data['questions'][0]
        self.assertEqual(question['options'], ['A check', 'A program'])
        self.assertNotIn('correct_index', question)
        self.assertNotIn('explanation', question)

    def test_empty_question_quiz_returns_clear_empty_state(self):
        self.quiz.questions = []
        with patch('academics.views._owned_topic_quiz', return_value=(self.quiz, self.subject, self.topic)):
            response = TopicQuizDetail().get(self.request, 'quiz-1')
        self.assertEqual(response.status_code, 404)
        self.assertIn('does not have any questions', response.data['detail'])

    def test_invalid_quiz_id_is_not_found(self):
        with patch('academics.views._owned_topic_quiz', return_value=(None, None, None)):
            response = TopicQuizDetail().get(self.request, 'invalid-id')
        self.assertEqual(response.status_code, 404)

    def test_backend_scores_and_returns_correct_and_incorrect_review(self):
        self.request.data = {'answers': [
            {'question_id': 'q1', 'selected_index': 0},
            {'question_id': 'q2', 'selected_index': 1},
        ]}
        with patch('academics.views._owned_topic_quiz', return_value=(self.quiz, self.subject, self.topic)), \
                patch.object(QuizAttempt, 'save', autospec=True, side_effect=lambda attempt: attempt):
            response = TopicQuizAttempts().post(self.request, 'quiz-1')

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['score'], 50)
        self.assertEqual(response.data['correct_count'], 1)
        self.assertEqual(response.data['incorrect_count'], 1)
        self.assertTrue(response.data['answers'][0]['is_correct'])
        self.assertFalse(response.data['answers'][1]['is_correct'])
        self.assertEqual(response.data['answers'][1]['correct_answer'], 'Unit')
        self.assertIn('test level', response.data['answers'][1]['explanation'])

    def test_missing_duplicate_foreign_or_invalid_answers_are_rejected(self):
        invalid_submissions = [
            {'answers': [{'question_id': 'q1', 'selected_index': 0}]},
            {'answers': [{'question_id': 'q1', 'selected_index': 0}, {'question_id': 'q1', 'selected_index': 1}]},
            {'answers': [{'question_id': 'q1', 'selected_index': 0}, {'question_id': 'foreign-q', 'selected_index': 0}]},
            {'answers': [{'question_id': 'q1', 'selected_index': 9}, {'question_id': 'q2', 'selected_index': 0}]},
            {'answers': [{'question_id': 'q1', 'selected_index': 0}, {'question_id': 'q2', 'selected_index': 0}], 'score': 100},
        ]
        with patch('academics.views._owned_topic_quiz', return_value=(self.quiz, self.subject, self.topic)), \
                patch.object(QuizAttempt, 'save', autospec=True) as save:
            for submission in invalid_submissions:
                self.request.data = submission
                response = TopicQuizAttempts().post(self.request, 'quiz-1')
                self.assertEqual(response.status_code, 400)
            save.assert_not_called()

    def test_database_failure_returns_503(self):
        with patch('academics.views._owned_topic_quiz', side_effect=ServerSelectionTimeoutError('quiz DB unavailable')):
            response = TopicQuizDetail().get(self.request, 'quiz-1')
        self.assertEqual(response.status_code, 503)

    def test_generate_requires_server_api_key(self):
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 2, 'syllabus_token': 'test-token',
        })
        cached = {'subject_id': 'subject-1', 'topics': [{'id': 'topic-1', 'title': 'Software Testing'}], 'text': 'Software Testing\n' + ('Test cases describe expected outcomes. ' * 5)}
        with patch.dict('os.environ', {}, clear=True), patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.cache.get', return_value=cached), patch('academics.views.TopicQuiz.objects') as quizzes:
            subjects.get.return_value = self.subject
            quizzes.return_value.first.return_value = None
            response = TopicQuizGenerate().post(request)
        self.assertEqual(response.status_code, 503)
        self.assertIn('OPENAI_API_KEY', response.data['detail'])

    @override_settings(QUIZ_AI_GENERATION_ENABLED=False)
    def test_disabled_generation_returns_coming_soon_without_openai_request(self):
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={})
        with patch('academics.views.requests.post') as openai:
            response = TopicQuizGenerate().post(request)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.data['error']['code'], 'feature_coming_soon')
        self.assertIn('under development', response.data['detail'])
        openai.assert_not_called()

    def test_generate_validates_topic_ownership_before_openai_call(self):
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['foreign-topic'], 'question_count': 2,
        })
        with patch.dict('os.environ', {'OPENAI_API_KEY': 'test-key'}), \
                patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.requests.post') as openai:
            subjects.get.return_value = self.subject
            response = TopicQuizGenerate().post(request)
        self.assertEqual(response.status_code, 400)
        openai.assert_not_called()

    def test_generate_rejects_invalid_question_count(self):
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 21,
        })
        with patch('academics.views.Subject.objects') as subjects:
            subjects.get.return_value = self.subject
            response = TopicQuizGenerate().post(request)
        self.assertEqual(response.status_code, 400)

    def test_question_bank_create_enforces_requested_count_and_topic_scope(self):
        from academics.views import TopicQuizCreateFromBank
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 1,
        })
        bank_query = MagicMock()
        bank_query.order_by.return_value = [self.quiz]
        bank_query.first.return_value = None
        with patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.TopicQuiz.objects', return_value=bank_query), \
                patch.object(TopicQuiz, 'save', autospec=True, side_effect=lambda quiz: setattr(quiz, 'id', 'bank-quiz') or quiz):
            subjects.get.return_value = self.subject
            response = TopicQuizCreateFromBank().post(request)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(len(response.data['questions']), 1)
        self.assertEqual(response.data['source'], 'question_bank')
        self.assertEqual(response.data['questions'][0]['topic_id'], 'topic-1')
        self.assertNotIn('correct_index', response.data['questions'][0])

    def test_question_bank_shortage_and_foreign_topic_are_rejected_without_saving(self):
        from academics.views import TopicQuizCreateFromBank
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 3,
        })
        bank_query = MagicMock()
        bank_query.order_by.return_value = [self.quiz]
        with patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.TopicQuiz.objects', return_value=bank_query), \
                patch.object(TopicQuiz, 'save', autospec=True) as save:
            subjects.get.return_value = self.subject
            response = TopicQuizCreateFromBank().post(request)
            self.assertEqual(response.status_code, 409)
            self.assertEqual(response.data['available_count'], 2)
            request.data['topic_ids'] = ['someone-elses-topic']
            invalid = TopicQuizCreateFromBank().post(request)
        self.assertEqual(invalid.status_code, 400)
        save.assert_not_called()

    def test_question_bank_availability_counts_unique_selected_topic_questions(self):
        from academics.views import TopicQuizAvailability
        self.quiz.questions.append(QuizQuestion(
            question_id='duplicate-q', text='  WHAT is a test case? ',
            options=['A check', 'A program'], correct_index=0,
        ))
        request = SimpleNamespace(
            user=SimpleNamespace(id=self.user_id),
            query_params={'subject_id': 'subject-1', 'topic_ids': 'topic-1'},
        )
        bank_query = MagicMock()
        bank_query.order_by.return_value = [self.quiz]
        with patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.TopicQuiz.objects', return_value=bank_query):
            subjects.get.return_value = self.subject
            response = TopicQuizAvailability().get(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['available_count'], 2)

    def test_question_bank_availability_and_creation_respect_difficulty(self):
        from academics.views import TopicQuizAvailability, TopicQuizCreateFromBank
        self.request.query_params['difficulty'] = 'hard'
        bank_query = MagicMock()
        bank_query.order_by.return_value = [self.quiz]
        with patch('academics.views.Subject.objects') as subjects, patch('academics.views.TopicQuiz.objects', return_value=bank_query):
            subjects.get.return_value = self.subject
            available = TopicQuizAvailability().get(self.request)
            created = TopicQuizCreateFromBank().post(SimpleNamespace(
                user=SimpleNamespace(id=self.user_id),
                data={'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 1, 'difficulty': 'hard'},
            ))
        self.assertEqual(available.status_code, 200)
        self.assertEqual(available.data['available_count'], 0)
        self.assertEqual(created.status_code, 409)
        self.assertEqual(created.data['available_count'], 0)

    def test_repeated_question_bank_creation_reuses_the_saved_selection(self):
        from academics.views import TopicQuizCreateFromBank
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 1,
        })
        bank_query = MagicMock()
        bank_query.order_by.return_value = [self.quiz]
        bank_query.first.return_value = None
        with patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.TopicQuiz.objects', return_value=bank_query), \
                patch.object(TopicQuiz, 'save', autospec=True, side_effect=lambda quiz: setattr(quiz, 'id', 'bank-quiz') or quiz) as save:
            subjects.get.return_value = self.subject
            first = TopicQuizCreateFromBank().post(request)
            saved_quiz = save.call_args.args[0]
            bank_query.first.return_value = saved_quiz
            second = TopicQuizCreateFromBank().post(request)
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(first.data['id'], second.data['id'])
        save.assert_called_once()

    def test_generate_saves_topic_scoped_questions_and_withholds_keys(self):
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 1, 'syllabus_token': 'test-token',
        })
        generated = {'questions': [{
            'topic_id': 'topic-1', 'question': 'What defines a test case?',
            'options': ['Input and expected output', 'A compiler', 'A test runner', 'A deployment step'], 'correct_index': 0,
            'explanation': 'The source describes test cases using expected results.',
        }]}
        fake_response = SimpleNamespace(
            raise_for_status=lambda: None,
            json=lambda: {'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': json.dumps(generated)}]}]},
        )
        cached = {'subject_id': 'subject-1', 'topics': [{'id': 'topic-1', 'title': 'Software Testing'}], 'text': 'Software Testing\n' + ('Test cases define inputs and expected results for software. ' * 4)}
        with patch.dict('os.environ', {'OPENAI_API_KEY': 'test-key'}), \
                patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.cache.get', return_value=cached), \
                patch('academics.views.TopicQuiz.objects') as quizzes, \
                patch('academics.views.requests.post', return_value=fake_response) as openai, \
                patch.object(TopicQuiz, 'save', autospec=True, side_effect=lambda quiz: setattr(quiz, 'id', 'generated-quiz') or quiz):
            subjects.get.return_value = self.subject
            quizzes.return_value.first.return_value = None
            response = TopicQuizGenerate().post(request)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['topic_ids'], ['topic-1'])
        self.assertEqual(response.data['source'], 'openai_generated')
        self.assertNotIn('correct_index', response.data['questions'][0])
        self.assertNotIn('explanation', response.data['questions'][0])
        self.assertEqual(openai.call_args.kwargs['json']['model'], 'gpt-4o-mini')
        payload = openai.call_args.kwargs['json']
        self.assertEqual(payload['input'][1]['content'] and json.loads(payload['input'][1]['content'])['requested_question_count'], 1)
        self.assertEqual(json.loads(payload['input'][1]['content'])['difficulty'], 'medium')
        self.assertIn('Software Testing', json.loads(payload['input'][1]['content'])['selected_topics'][0]['title'])

    def test_openai_provider_failures_return_specific_safe_error_codes(self):
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 1, 'syllabus_token': 'test-token',
        })
        cached = {'subject_id': 'subject-1', 'topics': [{'id': 'topic-1', 'title': 'Software Testing'}], 'text': 'Software Testing\n' + ('Testing uses defined expected results. ' * 5)}
        cases = [
            (401, {'code': 'invalid_api_key', 'type': 'authentication_error'}, 'invalid_credentials', 502),
            (429, {'code': 'insufficient_quota', 'type': 'insufficient_quota'}, 'insufficient_quota', 402),
            (429, {'code': 'rate_limit_exceeded', 'type': 'requests'}, 'rate_limit', 429),
            (400, {'code': 'model_not_found', 'param': 'model'}, 'model_configuration', 502),
        ]
        with patch.dict('os.environ', {'OPENAI_API_KEY': 'test-key'}), \
                patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.cache.get', return_value=cached), \
                patch('academics.views.TopicQuiz.objects') as quizzes:
            subjects.get.return_value = self.subject
            quizzes.return_value.first.return_value = None
            for status_code, error, expected_code, expected_status in cases:
                provider_response = SimpleNamespace(status_code=status_code, headers={'x-request-id': 'safe-id'}, json=lambda err=error: {'error': err})
                with patch('academics.views.requests.post', return_value=provider_response):
                    response = TopicQuizGenerate().post(request)
                self.assertEqual(response.status_code, expected_status)
                self.assertEqual(response.data['error']['code'], expected_code)
                self.assertNotIn('test-key', json.dumps(response.data))

    def test_connection_and_malformed_openai_responses_are_actionable(self):
        import requests
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 1, 'syllabus_token': 'test-token',
        })
        cached = {'subject_id': 'subject-1', 'topics': [{'id': 'topic-1', 'title': 'Software Testing'}], 'text': 'Software Testing\n' + ('Testing uses defined expected results. ' * 5)}
        with patch.dict('os.environ', {'OPENAI_API_KEY': 'test-key'}), \
                patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.cache.get', return_value=cached), \
                patch('academics.views.TopicQuiz.objects') as quizzes:
            subjects.get.return_value = self.subject
            quizzes.return_value.first.return_value = None
            with patch('academics.views.requests.post', side_effect=requests.ConnectionError('network unavailable')):
                unavailable = TopicQuizGenerate().post(request)
            malformed_response = SimpleNamespace(
                status_code=200, headers={'x-request-id': 'safe-id'},
                json=lambda: {'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': '{broken'}]}]},
            )
            with patch('academics.views.requests.post', return_value=malformed_response):
                malformed = TopicQuizGenerate().post(request)
        self.assertEqual(unavailable.status_code, 503)
        self.assertEqual(unavailable.data['error']['code'], 'provider_unavailable')
        self.assertEqual(malformed.status_code, 502)
        self.assertEqual(malformed.data['error']['code'], 'invalid_generated_response')

    def test_generated_quiz_database_failure_has_distinct_error(self):
        from academics.views import TopicQuizGenerate
        request = SimpleNamespace(user=SimpleNamespace(id=self.user_id), data={
            'subject_id': 'subject-1', 'topic_ids': ['topic-1'], 'question_count': 1, 'syllabus_token': 'test-token',
        })
        cached = {'subject_id': 'subject-1', 'topics': [{'id': 'topic-1', 'title': 'Software Testing'}], 'text': 'Software Testing\n' + ('Testing uses defined expected results. ' * 5)}
        generated = {'questions': [{'topic_id': 'topic-1', 'question': 'What defines a test case?', 'options': ['Inputs and expected output', 'A compiler', 'A database', 'A deployment'], 'correct_index': 0, 'explanation': 'The syllabus says this.'}]}
        provider_response = SimpleNamespace(status_code=200, headers={}, json=lambda: {'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': json.dumps(generated)}]}]})
        with patch.dict('os.environ', {'OPENAI_API_KEY': 'test-key'}), \
                patch('academics.views.Subject.objects') as subjects, \
                patch('academics.views.cache.get', return_value=cached), \
                patch('academics.views.TopicQuiz.objects') as quizzes, \
                patch('academics.views.requests.post', return_value=provider_response), \
                patch.object(TopicQuiz, 'save', autospec=True, side_effect=ServerSelectionTimeoutError('db unavailable')):
            subjects.get.return_value = self.subject
            quizzes.return_value.first.return_value = None
            response = TopicQuizGenerate().post(request)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.data['error']['code'], 'database_save_failed')


class QuizSyllabusUploadTests(SimpleTestCase):
    @staticmethod
    def syllabus_text():
        return (
            'Unit 1: Software Engineering\n'
            '1.1 Software Testing\n'
            'Software testing checks software against defined requirements and expected results. '
            'Test cases describe inputs, actions, and observable outcomes for a program.\n'
            '1.2 Testing Techniques\n'
            'Testing techniques include black box and white box approaches. '
            'Black box testing checks visible behavior; white box testing uses internal structure.\n'
        )

    def test_txt_bom_decodes_and_docx_extracts_paragraphs_and_tables(self):
        from academics import syllabus_processor
        with tempfile.TemporaryDirectory() as directory:
            text_path = Path(directory) / 'syllabus.txt'
            text_path.write_bytes(('\ufeff' + self.syllabus_text()).encode('utf-8'))
            extracted, warning = syllabus_processor.extract_text(str(text_path), '.txt')
            self.assertIsNone(warning)
            self.assertIn('Unit 1', extracted)
            self.assertNotIn('\ufeff', extracted)

            from docx import Document
            document = Document()
            document.add_paragraph('Chapter One: Testing')
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = 'Input'
            table.rows[0].cells[1].text = 'Expected output'
            docx_path = Path(directory) / 'syllabus.docx'
            document.save(docx_path)
            extracted_docx, warning = syllabus_processor.extract_text(str(docx_path), '.docx')
        self.assertIsNone(warning)
        self.assertIn('Chapter One: Testing', extracted_docx)
        self.assertIn('Input | Expected output', extracted_docx)

    def test_pdf_scanned_page_reports_ocr_required(self):
        from academics import syllabus_processor
        from pypdf import PdfWriter
        buffer = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(buffer)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'scan.pdf'
            path.write_bytes(buffer.getvalue())
            extracted, warning = syllabus_processor.extract_text(str(path), '.pdf')
        self.assertEqual(extracted.strip(), '')
        self.assertIn('scanned', warning.lower())

    def test_upload_validation_checks_extension_signature_content_and_size(self):
        from academics import syllabus_processor
        unsupported = SimpleUploadedFile('notes.exe', b'content', content_type='text/plain')
        fake_pdf = SimpleUploadedFile('notes.pdf', b'not a pdf', content_type='application/pdf')
        fake_docx = SimpleUploadedFile('notes.docx', b'not a zip', content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        empty = SimpleUploadedFile('empty.txt', b'', content_type='text/plain')
        oversized = SimpleUploadedFile('large.txt', b'12345', content_type='text/plain')
        for uploaded in (unsupported, fake_pdf, fake_docx, empty):
            valid, message = syllabus_processor.validate_quiz_upload(uploaded)
            self.assertFalse(valid)
            self.assertTrue(message)
        self.assertEqual(syllabus_processor.validate_quiz_upload(oversized, max_file_size=4)[0], False)
        # Browser MIME declarations are not trusted; extension and actual text encoding are checked.
        valid_txt = SimpleUploadedFile('valid.txt', self.syllabus_text().encode('utf-8'), content_type='application/pdf')
        self.assertEqual(syllabus_processor.validate_quiz_upload(valid_txt), (True, None))

    def test_upload_extracts_topics_without_mutating_saved_subject(self):
        from academics.views import QuizSyllabusUpload
        user_id = 'upload-student'
        subject = SimpleNamespace(id='subject-upload', name='Software Engineering', topics=[], syllabus_files=[])
        uploaded = SimpleUploadedFile('semester.txt', self.syllabus_text().encode('utf-8'), content_type='application/octet-stream')
        request = SimpleNamespace(
            user=SimpleNamespace(id=user_id), data={'subject_id': 'subject-upload'}, FILES={'file': uploaded},
        )
        with patch('academics.views.Subject.objects') as subjects, patch('academics.views.cache.set') as cache_set:
            subjects.get.return_value = subject
            response = QuizSyllabusUpload().post(request)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['file_type'], 'TXT')
        self.assertEqual(response.data['extraction_status'], 'complete')
        self.assertGreaterEqual(response.data['total_found'], 2)
        self.assertIn('Software Testing', [topic['title'] for topic in response.data['topics']])
        self.assertIn('Software testing checks', response.data['preview'])
        self.assertEqual(subject.topics, [])
        self.assertEqual(subject.syllabus_files, [])
        self.assertTrue(cache_set.call_args.args[0].startswith(f'quiz-syllabus:{user_id}:'))
        self.assertEqual(cache_set.call_args.kwargs['timeout'], 30 * 60)

    def test_upload_enforces_configured_server_size_limit(self):
        from academics.views import QuizSyllabusUpload
        subject = SimpleNamespace(id='subject-upload', name='Software Engineering', topics=[], syllabus_files=[])
        uploaded = SimpleUploadedFile('large.txt', self.syllabus_text().encode('utf-8'), content_type='text/plain')
        request = SimpleNamespace(
            user=SimpleNamespace(id='upload-student'), data={'subject_id': 'subject-upload'}, FILES={'file': uploaded},
        )
        with override_settings(QUIZ_SYLLABUS_MAX_FILE_SIZE=8), patch('academics.views.Subject.objects') as subjects:
            subjects.get.return_value = subject
            response = QuizSyllabusUpload().post(request)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data['code'], 'invalid_syllabus_file')

    def test_uploaded_syllabus_subject_is_owner_scoped(self):
        from academics.views import QuizSyllabusUpload
        request = SimpleNamespace(
            user=SimpleNamespace(id='other-student'), data={'subject_id': 'private-subject'}, FILES={},
        )
        with patch('academics.views.Subject.objects') as subjects:
            subjects.get.side_effect = DoesNotExist
            response = QuizSyllabusUpload().post(request)
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.data['code'], 'subject_not_found')


class StudyPlanSaveTests(SimpleTestCase):
    def test_duplicate_session_ids_are_rejected(self):
        request = SimpleNamespace(
            user=SimpleNamespace(id='student-1'),
            data={
                'title': 'Study',
                'start_date': '2026-10-29',
                'end_date': '2026-10-29',
                'sessions': [
                    {'session_id': 'duplicate'},
                    {'session_id': 'duplicate'},
                ],
            },
        )

        response = StudyPlanListCreate().post(request)

        self.assertEqual(response.status_code, 400)

    @patch('academics.views.StudyPlan.objects')
    def test_retrying_same_preview_returns_existing_plan(self, plans):
        existing = SimpleNamespace(
            id='plan-1',
            title='Study',
            start_date='2026-10-29',
            end_date='2026-10-29',
            daily_hours='2',
            created_at=date(2026, 10, 29),
            sessions=[{'session_id': 'session-1', 'topic_id': 'topic-1'}],
        )
        plans.return_value = [existing]
        request = SimpleNamespace(
            user=SimpleNamespace(id='student-1'),
            data={
                'title': 'Study',
                'start_date': '2026-10-29',
                'end_date': '2026-10-29',
                'daily_hours': '2',
                'sessions': [{'session_id': 'session-1', 'topic_id': 'topic-1'}],
            },
        )

        response = StudyPlanListCreate().post(request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['id'], 'plan-1')
        plans.assert_called_once_with(
            user_id='student-1',
            title='Study',
            start_date='2026-10-29',
            end_date='2026-10-29',
            daily_hours='2',
        )


class StudyPlannerDatabaseFailureTests(SimpleTestCase):
    def setUp(self):
        self.request = SimpleNamespace(
            user=SimpleNamespace(id='student-1'),
            data={'start_date': '2026-10-29', 'end_date': '2026-10-29', 'daily_hours': 2},
        )

    @patch('academics.views.StudyPlan.objects')
    def test_plan_list_returns_503_when_mongo_is_unavailable(self, plans):
        plans.side_effect = ServerSelectionTimeoutError('test database unavailable')

        response = StudyPlanListCreate().get(self.request)

        self.assertEqual(response.status_code, 503)
        self.assertIn('temporarily unavailable', response.data['detail'])

    @patch('academics.views.Subject.objects')
    def test_generation_returns_503_when_mongo_is_unavailable(self, subjects):
        subjects.side_effect = ServerSelectionTimeoutError('test database unavailable')

        response = GenerateStudyPlan().post(self.request)

        self.assertEqual(response.status_code, 503)

    @patch('academics.views.StudyPlan.objects')
    def test_detail_does_not_misreport_database_failure_as_not_found(self, plans):
        plans.get.side_effect = ServerSelectionTimeoutError('test database unavailable')

        response = StudyPlanDetail().get(self.request, 'plan-id')

        self.assertEqual(response.status_code, 503)


@unittest.skipUnless(
    os.environ.get('STUDY_PLANNER_LIVE_MONGO') == '1',
    'Set STUDY_PLANNER_LIVE_MONGO=1 and DB_NAME to a disposable test database to run MongoDB integration checks.',
)
class LiveMongoStudyPlannerTests(SimpleTestCase):
    """Opt-in integration tests against an explicitly named disposable MongoDB database."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.db = get_db()
        cls.db_name = cls.db.name
        if not cls.db_name.startswith('edusphere_studyplanner_test_'):
            raise RuntimeError('Refusing MongoDB integration tests outside the isolated study planner test database.')
        cls.db.client.admin.command('ping')
        cls.db.client.drop_database(cls.db_name)

    @classmethod
    def tearDownClass(cls):
        try:
            cls.db.client.drop_database(cls.db_name)
        finally:
            super().tearDownClass()

    def setUp(self):
        self.user_id = f'student-{uuid4()}'
        self.other_user_id = f'student-{uuid4()}'
        self.subject = Subject(
            user_id=self.user_id,
            name='Planner Integration Subject',
            topics=[SyllabusTopic(
                id='topic-1', title='Persisted topic', difficulty='medium',
                status='in_progress', revision_required=True,
            )],
        ).save()

    @staticmethod
    def request(user_id, data=None):
        return SimpleNamespace(user=SimpleNamespace(id=user_id), data=data or {})

    def test_generated_preview_saves_reloads_and_is_owner_scoped(self):
        preview_response = GenerateStudyPlan().post(self.request(self.user_id, {
            'start_date': '2026-10-29',
            'end_date': '2026-10-29',
            'daily_hours': 2,
            'subject_ids': [str(self.subject.id)],
        }))
        self.assertEqual(preview_response.status_code, 200)
        sessions = preview_response.data['sessions']
        self.assertEqual(len(sessions), 1)
        self.assertAlmostEqual(sessions[0]['priority_score'], 0.392, places=3)

        plan_data = {
            'title': 'Persisted Plan',
            'start_date': '2026-10-29',
            'end_date': '2026-10-29',
            'daily_hours': '2',
            'sessions': sessions,
        }
        create_response = StudyPlanListCreate().post(self.request(self.user_id, plan_data))
        self.assertEqual(create_response.status_code, 201)

        # A fresh API request, as after a page refresh or new login, reads the saved record.
        reloaded = StudyPlanListCreate().get(self.request(self.user_id))
        self.assertEqual(len(reloaded.data), 1)
        self.assertEqual(reloaded.data[0]['id'], create_response.data['id'])
        self.assertEqual(reloaded.data[0]['sessions'][0]['topic_id'], 'topic-1')
        persisted_topic = Subject.objects.get(id=self.subject.id, user_id=self.user_id).topics[0]
        self.assertEqual(persisted_topic.status, 'in_progress')
        self.assertTrue(persisted_topic.revision_required)

        retry_response = StudyPlanListCreate().post(self.request(self.user_id, plan_data))
        self.assertEqual(retry_response.status_code, 200)
        self.assertEqual(StudyPlan.objects(user_id=self.user_id).count(), 1)
        self.assertEqual(StudyPlanListCreate().get(self.request(self.other_user_id)).data, [])

    def test_session_completion_updates_database_and_insights_once(self):
        plan = StudyPlan(
            user_id=self.user_id,
            title='Completion Plan',
            start_date=date.today().isoformat(),
            end_date=date.today().isoformat(),
            daily_hours='2',
            sessions=[{
                'session_id': 'session-1',
                'subject': self.subject.name,
                'subject_id': str(self.subject.id),
                'topic_id': 'topic-1',
                'topic': 'Persisted topic',
                'date': date.today().isoformat(),
                'duration': '30 min',
                'duration_minutes': 30,
                'is_completed': False,
            }],
        ).save()
        request = self.request(self.user_id, {'session_id': 'session-1', 'is_completed': True})

        first = StudyPlanDetail().patch(request, str(plan.id))
        second = StudyPlanDetail().patch(request, str(plan.id))

        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        persisted_subject = Subject.objects.get(id=self.subject.id, user_id=self.user_id)
        persisted_plan = StudyPlan.objects.get(id=plan.id, user_id=self.user_id)
        self.assertTrue(persisted_subject.topics[0].is_completed)
        self.assertEqual(persisted_subject.topics[0].status, 'completed')
        self.assertTrue(persisted_subject.topics[0].revision_required)
        self.assertTrue(persisted_plan.sessions[0]['is_completed'])

        insights = StudyInsights().get(self.request(self.user_id))
        self.assertEqual(insights.data['completed_topics'], 1)
        self.assertEqual(insights.data['completed_sessions'], 1)
        self.assertEqual(insights.data['total_sessions'], 1)

        foreign_plan_response = StudyPlanDetail().get(self.request(self.other_user_id), str(plan.id))
        self.assertEqual(foreign_plan_response.status_code, 404)
        foreign_patch_response = StudyPlanDetail().patch(
            self.request(self.other_user_id, {'session_id': 'session-1', 'is_completed': False}),
            str(plan.id),
        )
        self.assertEqual(foreign_patch_response.status_code, 404)
        self.assertTrue(Subject.objects.get(id=self.subject.id, user_id=self.user_id).topics[0].is_completed)

    def test_academic_assistant_reads_saved_records_and_isolates_students(self):
        user = User(
            username=f'assistant-{uuid4()}@example.test',
            email=f'assistant-{uuid4()}@example.test',
        ).save()
        other_user = User(
            username=f'assistant-{uuid4()}@example.test',
            email=f'assistant-{uuid4()}@example.test',
        ).save()
        user_id = str(user.id)
        other_user_id = str(other_user.id)
        own_subject = Subject(user_id=user_id, name='My subject').save()
        other_subject = Subject(
            user_id=other_user_id,
            name='Another student subject',
        ).save()
        Assignment(
            user_id=user_id,
            subject_id=str(own_subject.id),
            title='My pending assignment',
            due_date=(date.today() + timedelta(days=2)).isoformat(),
            priority='high',
        ).save()
        Assignment(
            user_id=other_user_id,
            subject_id=str(other_subject.id),
            title='Private foreign assignment',
            due_date=(date.today() + timedelta(days=1)).isoformat(),
            priority='high',
        ).save()

        client = APIClient()
        unauthorized = client.post('/api/academics/assistant/ask/', {'question': 'What is my progress?'}, format='json')
        self.assertEqual(unauthorized.status_code, 403)
        access_token, _ = make_tokens(user)
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {access_token}')
        response = client.post('/api/academics/assistant/ask/', {
            'question': 'Which assignments are pending?'
        }, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertIn('My pending assignment', response.data['answer'])
        self.assertNotIn('Private foreign assignment', response.data['answer'])
        other_access_token, _ = make_tokens(other_user)
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {other_access_token}')
        other_response = client.post('/api/academics/assistant/ask/', {
            'question': 'Which assignments are pending?'
        }, format='json')
        self.assertEqual(other_response.status_code, 200)
        self.assertIn('Private foreign assignment', other_response.data['answer'])
        self.assertNotIn('My pending assignment', other_response.data['answer'])

    def test_topic_quiz_import_authenticated_attempts_persist_and_isolate(self):
        user = User(
            username=f'quiz-{uuid4()}@example.test',
            email=f'quiz-{uuid4()}@example.test',
        ).save()
        other_user = User(
            username=f'quiz-{uuid4()}@example.test',
            email=f'quiz-{uuid4()}@example.test',
        ).save()
        user_id = str(user.id)
        other_user_id = str(other_user.id)
        subject = Subject(
            user_id=user_id,
            name='Software Engineering',
            topics=[
                SyllabusTopic(id='testing-topic', title='Software Testing'),
                SyllabusTopic(id='design-topic', title='Software Design'),
            ],
        ).save()
        with tempfile.TemporaryDirectory() as temp_dir:
            question_file = Path(temp_dir) / 'testing_quiz.json'
            question_file.write_text(json.dumps({
                'title': 'Testing Fundamentals',
                'questions': [
                    {'question': 'What does a test case define?', 'options': ['Input and expected output', 'Only source code'], 'correct_index': 0, 'explanation': 'A test case specifies an input and expected result.'},
                    {'question': 'Which is a test level?', 'options': ['Unit testing', 'Linking'], 'correct_index': 0, 'explanation': 'Unit testing is a recognized testing level.'},
                ],
            }), encoding='utf-8')
            call_command(
                'seed_topic_quiz', user_id=user_id, subject_id=str(subject.id),
                topic_id='testing-topic', file=str(question_file), stdout=StringIO(),
            )

        quiz = TopicQuiz.objects.get(user_id=user_id, subject_id=str(subject.id), topic_id='testing-topic')
        self.assertEqual(quiz.source, 'author_created')
        self.assertEqual(len(quiz.questions), 2)

        client = APIClient()
        unauthenticated = client.get(f'/api/academics/quizzes/{quiz.id}/')
        self.assertEqual(unauthenticated.status_code, 403)
        access_token, _ = make_tokens(user)
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {access_token}')
        quiz_list = client.get('/api/academics/quizzes/', {
            'subject_id': str(subject.id), 'topic_id': 'testing-topic',
        })
        self.assertEqual(quiz_list.status_code, 200)
        self.assertEqual(len(quiz_list.data), 1)
        self.assertEqual(quiz_list.data[0]['subject_name'], 'Software Engineering')
        self.assertEqual(quiz_list.data[0]['topic_name'], 'Software Testing')
        unrelated_topic_list = client.get('/api/academics/quizzes/', {
            'subject_id': str(subject.id), 'topic_id': 'design-topic',
        })
        self.assertEqual(unrelated_topic_list.status_code, 200)
        self.assertEqual(unrelated_topic_list.data, [])

        detail = client.get(f'/api/academics/quizzes/{quiz.id}/')
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(len(detail.data['questions']), 2)
        self.assertNotIn('correct_index', detail.data['questions'][0])
        self.assertNotIn('explanation', detail.data['questions'][0])

        answer_payload = {'answers': [
            {'question_id': quiz.questions[0].question_id, 'selected_index': 0},
            {'question_id': quiz.questions[1].question_id, 'selected_index': 1},
        ]}
        manipulated = client.post(
            f'/api/academics/quizzes/{quiz.id}/attempts/',
            {**answer_payload, 'score': 100}, format='json',
        )
        self.assertEqual(manipulated.status_code, 400)
        first_attempt = client.post(f'/api/academics/quizzes/{quiz.id}/attempts/', answer_payload, format='json')
        second_attempt = client.post(f'/api/academics/quizzes/{quiz.id}/attempts/', answer_payload, format='json')
        self.assertEqual((first_attempt.status_code, second_attempt.status_code), (201, 201))
        self.assertEqual(first_attempt.data['score'], 50)
        self.assertEqual(second_attempt.data['score'], first_attempt.data['score'])
        self.assertEqual(first_attempt.data['correct_count'], 1)
        self.assertEqual(first_attempt.data['incorrect_count'], 1)
        self.assertIn('correct_index', first_attempt.data['answers'][0])
        self.assertIn('recognized testing level', first_attempt.data['answers'][1]['explanation'])

        history = client.get(f'/api/academics/quizzes/{quiz.id}/attempts/')
        self.assertEqual(history.status_code, 200)
        self.assertEqual(len(history.data), 2)
        self.assertEqual(QuizAttempt.objects(user_id=user_id, quiz_id=str(quiz.id)).count(), 2)

        other_access, _ = make_tokens(other_user)
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {other_access}')
        self.assertEqual(client.get(f'/api/academics/quizzes/{quiz.id}/').status_code, 404)
        self.assertEqual(client.get(f'/api/academics/quizzes/{quiz.id}/attempts/').status_code, 404)
        self.assertEqual(client.post(f'/api/academics/quizzes/{quiz.id}/attempts/', answer_payload, format='json').status_code, 404)
        self.assertEqual(QuizAttempt.objects(user_id=other_user_id, quiz_id=str(quiz.id)).count(), 0)

        client.credentials(HTTP_AUTHORIZATION=f'Bearer {access_token}')
        bank_payload = {'subject_id': str(subject.id), 'topic_ids': ['testing-topic'], 'question_count': 2}
        assembled = client.post('/api/academics/quizzes/create/', bank_payload, format='json')
        repeated = client.post('/api/academics/quizzes/create/', bank_payload, format='json')
        self.assertEqual(assembled.status_code, 201, assembled.data)
        self.assertEqual(assembled.data['source'], 'question_bank')
        self.assertEqual(repeated.status_code, 200)
        self.assertEqual(assembled.data['id'], repeated.data['id'])
        self.assertEqual(TopicQuiz.objects(user_id=user_id, source='question_bank').count(), 1)
        self.assertEqual(len(assembled.data['questions']), 2)
        self.assertEqual(assembled.data['topic_ids'], ['testing-topic'])

    def test_generated_quiz_persists_and_is_retrievable_only_by_owner(self):
        user = User(username=f'generated-quiz-{uuid4()}@example.test', email=f'generated-quiz-{uuid4()}@example.test').save()
        other = User(username=f'generated-quiz-{uuid4()}@example.test', email=f'generated-quiz-{uuid4()}@example.test').save()
        user_id = str(user.id)
        subject = Subject(
            user_id=user_id, name='Data Structures', topics=[], syllabus_files=[],
        ).save()
        client = APIClient()
        access, _ = make_tokens(user)
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {access}')
        unauthenticated_upload = APIClient().post('/api/academics/quizzes/syllabus/', {
            'subject_id': str(subject.id),
            'file': SimpleUploadedFile('no-auth.txt', b'private text', content_type='text/plain'),
        })
        self.assertEqual(unauthenticated_upload.status_code, 403)
        upload_text = (
            'Unit 1: Data Structures\n'
            '1.1 Binary Trees\n'
            + ('Binary trees contain nodes with at most two children. This structure supports hierarchical data and traversal. ' * 3)
            + '\n1.2 Complexity Analysis\n'
            + ('Time complexity describes how algorithm runtime changes as input size increases. It helps compare algorithms. ' * 3)
        )
        upload_response = client.post('/api/academics/quizzes/syllabus/', {
            'subject_id': str(subject.id),
            'file': SimpleUploadedFile('structures.txt', upload_text.encode('utf-8'), content_type='text/plain'),
        })
        self.assertEqual(upload_response.status_code, 200, upload_response.data)
        self.assertEqual([topic['title'] for topic in upload_response.data['topics']], ['Binary Trees', 'Complexity Analysis'])
        topic_ids = [topic['id'] for topic in upload_response.data['topics']]
        generated = {'questions': [
            {
                'topic_id': topic_ids[0], 'question': 'Which property describes a binary tree?',
                'options': ['Each node has at most two children', 'Each node has exactly four children', 'Nodes have no links', 'It is always balanced'],
                'correct_index': 0, 'explanation': 'The syllabus describes the binary tree child limit.',
            },
            {
                'topic_id': topic_ids[1], 'question': 'What does time complexity describe?',
                'options': ['How runtime grows with input', 'The color of a node', 'The number of source files', 'A memory address'],
                'correct_index': 0, 'explanation': 'The syllabus relates complexity to input size and runtime.',
            },
        ]}
        fake_response = SimpleNamespace(
            status_code=200, headers={'x-request-id': 'integration-test-request'},
            json=lambda: {'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': json.dumps(generated)}]}]},
        )
        generate_payload = {
            'subject_id': str(subject.id), 'topic_ids': topic_ids, 'question_count': 2,
            'difficulty': 'hard', 'syllabus_token': upload_response.data['upload_token'],
        }
        with patch.dict(os.environ, {'OPENAI_API_KEY': 'integration-test-key'}), \
                patch('academics.views.requests.post', return_value=fake_response) as openai:
            generated_response = client.post('/api/academics/quizzes/generate/', generate_payload, format='json')
            repeated_response = client.post('/api/academics/quizzes/generate/', generate_payload, format='json')
        self.assertEqual(generated_response.status_code, 201, generated_response.data)
        self.assertEqual(repeated_response.status_code, 200)
        self.assertEqual(generated_response.data['id'], repeated_response.data['id'])
        openai.assert_called_once()
        quiz_id = generated_response.data['id']
        stored = TopicQuiz.objects.get(id=quiz_id, user_id=user_id)
        self.assertEqual(stored.source, 'openai_generated')
        self.assertEqual(stored.topic_ids, topic_ids)
        self.assertEqual(stored.topic_names, ['Binary Trees', 'Complexity Analysis'])
        self.assertEqual(stored.difficulty, 'hard')
        self.assertEqual(subject.topics, [])  # Uploaded topics were not imported into the saved subject.
        self.assertEqual(stored.questions[0].correct_index, 0)
        self.assertEqual(stored.questions[0].difficulty, 'hard')
        self.assertNotIn('correct_index', generated_response.data['questions'][0])
        detail = client.get(f'/api/academics/quizzes/{quiz_id}/')
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.data['questions'][0]['question_id'], stored.questions[0].question_id)
        saved_list = client.get('/api/academics/quizzes/', {
            'subject_id': str(subject.id), 'topic_id': topic_ids[0], 'upload_token': upload_response.data['upload_token'],
        })
        self.assertEqual(saved_list.status_code, 200)
        self.assertEqual(len(saved_list.data), 1)
        generated_attempt = client.post(f'/api/academics/quizzes/{quiz_id}/attempts/', {
            'answers': [{'question_id': question['question_id'], 'selected_index': 0} for question in generated_response.data['questions']],
        }, format='json')
        self.assertEqual(generated_attempt.status_code, 201)
        self.assertEqual(generated_attempt.data['score'], 100)
        self.assertEqual(len(client.get(f'/api/academics/quizzes/{quiz_id}/attempts/').data), 1)
        other_access, _ = make_tokens(other)
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {other_access}')
        self.assertEqual(client.get(f'/api/academics/quizzes/{quiz_id}/').status_code, 404)
        self.assertEqual(client.post('/api/academics/quizzes/syllabus/', {
            'subject_id': str(subject.id), 'file': SimpleUploadedFile('private.txt', upload_text.encode('utf-8'), content_type='text/plain'),
        }).status_code, 404)
